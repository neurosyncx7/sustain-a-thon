from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from notes import SCRIPT, QA, RECORDING
NAVY=RGBColor(0x1B,0x3A,0x7A); MUTED=RGBColor(0x5B,0x6B,0x8C)
d=Document()
for s in d.sections: s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(2)
st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(12)
def h(t,lvl):
    p=d.add_heading(t,lvl)
    for r in p.runs: r.font.color.rgb=NAVY; r.font.name='Arial'
h("Vāyu Lekha · Presentation script",0)
p=d.add_paragraph(); r=p.add_run("Team Prahari · ST-911 · PS-13-S3 · Sustain-a-thon 2026"); r.bold=True
d.add_paragraph("Read the plain text out loud. Words in [brackets] are actions, not lines. Talk to the jury, not the screen. "
                "If a round is short, keep the first two paragraphs of each slide and skip the rest; the slide still carries the numbers.")
t=d.add_table(rows=1,cols=4); t.style='Light Grid Accent 1'
for c,x in zip(t.rows[0].cells,["#","Slide","Speaker","Time"]): c.text=x
for i,s in enumerate(SCRIPT,1):
    row=t.add_row().cells; row[0].text=str(i); row[1].text=s['title']; row[2].text=s['speaker']; row[3].text=s['time']
d.add_paragraph("Total: about 9 to 10 minutes with the live demo. Change the speaker split freely; the lines work for anyone.").runs[0].italic=True
for i,s in enumerate(SCRIPT,1):
    d.add_page_break()
    h(f"Slide {i} · {s['title']}",1)
    p=d.add_paragraph(); r=p.add_run(f"{s['speaker']}  ·  about {s['time']}"); r.bold=True; r.font.color.rgb=MUTED
    for para in s['say']:
        q=d.add_paragraph(para); q.paragraph_format.space_after=Pt(8); q.paragraph_format.line_spacing=1.25
        for r in q.runs: r.font.size=Pt(13)
    h("On screen",3)
    for c in s['show']:
        q=d.add_paragraph(c, style='List Bullet')
        for r in q.runs: r.italic=True
d.add_page_break(); h("Likely jury questions",1)
for qq,a in QA:
    p=d.add_paragraph(); r=p.add_run("Q. "+qq); r.bold=True; r.font.color.rgb=NAVY
    d.add_paragraph(a).paragraph_format.space_after=Pt(10)
d.add_page_break(); h("Backup screen recording: shot list",1)
for x in RECORDING: d.add_paragraph(x, style='List Number')
h("Sources for the numbers on the slides",2)
for x in ["Methane ~80× CO₂ over 20 years: IPCC AR6 WG1, Table 7.15 (GWP20 82.5 fossil, 79.7 non-fossil).",
          "India ~30 Mt CH₄ in 2023, second after China; 2,438 legacy dumpsites; waste 20 → 76 Mt CO₂e by 2030: Harvard Salata Institute, Research Brief 6 (Nov 2024), citing IEA and MoEFCC.",
          "~30% of warming since the Industrial Revolution: IEA Global Methane Tracker 2025.",
          "Sentinel-5P orbit repeats in 16 days / 227 orbits: Copernicus SentiWiki, S5P mission page.",
          "Carbon Mapper 2025, Hyderabad (Secunderabad) and Mumbai landfills in the world's top 25: Down To Earth, 22 April 2026.",
          "Everything else (pixels, files, rates, z, q, ablation, OBC, VOIT, WRPI): our pipeline, data-pipeline/ in github.com/neurosyncx7/sustain-a-thon, inventory run of 26 Sep 2026 23:16 UTC."]:
    d.add_paragraph(x, style='List Bullet')
d.save("Prahari_ST-911_Presentation_Script.docx"); print("ok")
